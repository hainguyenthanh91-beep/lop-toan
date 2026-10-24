"""Kiểm tra file Excel mẫu do bộ ghi Dart tạo ra (chạy trong CI)."""
import sys, zipfile
import openpyxl

f = sys.argv[1]
z = zipfile.ZipFile(f)
assert z.testzip() is None
wb = openpyxl.load_workbook(f)
names = wb.sheetnames
ws = wb[names[0]]
assert ws["A1"].value.startswith("ĐIỂM DANH"), ws["A1"].value
assert ws["C5"].value == "✓" and ws["C5"].fill.fgColor.rgb == "FFDCFCE7", (ws["C5"].value, ws["C5"].fill.fgColor.rgb)
assert ws["E5"].value == 150000 and ws["E5"].number_format == "#,##0", ws["E5"].number_format
assert ws["C6"].value == 8.5 and ws["C6"].number_format == "0.0"
assert ws.freeze_panes == "C5", ws.freeze_panes
assert "A1:G1" in [str(r) for r in ws.merged_cells.ranges]
assert ws["B2"].value.endswith('"test" & <xml>')
assert len(names) == 3 and names[2] == "Học phí (2)", names
print(f"::notice title=xlsx OK::sheets={names} freeze={ws.freeze_panes}")
